"""Build a publication slide (.pptx) summarizing the datasets that feed ADME model training.

One 16:9 slide: title + a table (dataset, nature, source, unique compounds, endpoints, role).
Counts are computed from the cached parquets (aggregate only — no SMILES/structures). SERAC palette.
Run (env ML):  python python/make_datasets_slide.py  ->  output/ppt/adme_datasets.pptx
"""
import glob, os
from pathlib import Path
import pandas as pd, yaml
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from ppt_style import horizontal_rules, plain_table, NAVY, HEADER_BG, BORDER

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / 'autoresearch/predict_adme'
OUT = ROOT / 'output/ppt'; OUT.mkdir(parents=True, exist_ok=True)
ENDPOINTS = ['solubility', 'logd', 'hlm', 'mlm', 'rlm', 'caco2', 'mdck', 'ppb']
PALETTE = yaml.safe_load((ROOT / 'config/config.yaml').read_text())['SERAC_C']


def _hex(h):
    return RGBColor.from_string(h.lstrip('#'))


def _uniq(files, col='smiles'):
    s = set()
    for f in files:
        s |= set(pd.read_parquet(f, columns=[col])[col].dropna())
    return len(s)


def _eps_for(prefix):
    return [os.path.basename(f).replace(prefix, '').replace('.parquet', '') for f in glob.glob(f'{CACHE}/{prefix}*.parquet')]


def gather_rows():
    """Return the data rows for the table (aggregate counts computed from the parquets)."""
    tgt = pd.read_parquet(CACHE / 'internal_targets.parquet')
    exp_files = [CACHE / f'public_{e}.parquet' for e in ENDPOINTS if (CACHE / f'public_{e}.parquet').exists()]
    exp_eps = [e for e in ENDPOINTS if (CACHE / f'public_{e}.parquet').exists()]
    nvs = glob.glob(f'{CACHE}/public_novartis_*.parquet')
    adm = glob.glob(f'{CACHE}/public_admetlab_*.parquet')
    order = {e: i for i, e in enumerate(ENDPOINTS)}
    fmt = lambda eps: ', '.join(sorted(set(eps), key=lambda e: order[e]))
    return [
        ['Internal — SERAC (bRo5: PROTACs, molecular glues)', 'Experimental (in-house)', 'CDD Vault',
         f'{len(tgt):,}', 'all 8', 'Prediction target / fine-tune'],
        ['Experimental public', 'Experimental', 'TDC · Biogen-Fang · harmonized solubility',
         f'{_uniq(exp_files):,}', fmt(exp_eps), 'Augment (real labels)'],
        ['Novartis / NIBR', 'Predicted (in-silico)', 'ProtacDB2.0 · ZINC · ChEMBL',
         f'{_uniq(nvs):,}', fmt(_eps_for("public_novartis_")), 'Augment (pseudo-labels)'],
        ['ADMETlab', 'Predicted (in-silico)', 'PROTAC-DB patent compounds',
         f'{_uniq(adm):,}', fmt(_eps_for("public_admetlab_")), 'Augment (pseudo-labels, selective)'],
    ]


def build(rows):
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])            # blank

    # navy title + subtitle (light/navy scheme — matches the potency table)
    title = slide.shapes.add_textbox(Inches(0.29), Inches(0.35), Inches(12.7), Inches(0.7)).text_frame
    title.text = 'Datasets used to train the ADME models'
    title.paragraphs[0].font.size = Pt(28); title.paragraphs[0].font.bold = True; title.paragraphs[0].font.color.rgb = _hex(NAVY)

    sub = slide.shapes.add_textbox(Inches(0.31), Inches(1.05), Inches(12.6), Inches(0.5)).text_frame
    sub.word_wrap = True; sub.text = ('325 in-house bRo5 measurements augmented with public data; '
                                      'predicted sources (Novartis, ADMETlab) used selectively per endpoint.')
    sub.paragraphs[0].font.size = Pt(13); sub.paragraphs[0].font.italic = True; sub.paragraphs[0].font.color.rgb = _hex('#555555')

    header = ['Dataset', 'Nature', 'Source', 'Unique cmpds', 'Endpoints covered', 'Role in training']
    widths = [2.9, 1.7, 2.6, 1.2, 2.75, 1.6]           # sum 12.75 -> fits the 13.33" slide
    tbl = slide.shapes.add_table(len(rows) + 1, len(header), Inches(0.29), Inches(1.85),
                                 Inches(sum(widths)), Inches(3.9)).table
    plain_table(tbl)                                   # no style borders/banding — horizontal rules only
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for j, h in enumerate(header):                                # header row: light bg, navy bold
        c = tbl.cell(0, j); c.text = h; c.fill.solid(); c.fill.fore_color.rgb = _hex(HEADER_BG)
        c.margin_left = c.margin_right = Inches(0.12); c.margin_top = c.margin_bottom = Inches(0.03)
        pr = c.text_frame.paragraphs[0]; pr.font.bold = True; pr.font.size = Pt(12); pr.font.color.rgb = _hex(NAVY)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE; horizontal_rules(c, BORDER)
    for i, row in enumerate(rows, start=1):                        # data rows: white bg, navy text
        for j, val in enumerate(row):
            c = tbl.cell(i, j); c.text = str(val); c.fill.solid(); c.fill.fore_color.rgb = _hex('#FFFFFF')
            c.margin_left = c.margin_right = Inches(0.12); c.margin_top = c.margin_bottom = Inches(0.03)
            pr = c.text_frame.paragraphs[0]; pr.font.size = Pt(10.5); pr.font.color.rgb = _hex(NAVY)
            pr.font.bold = (j in (0, 3))                          # bold the dataset name + the compound count
            c.vertical_anchor = MSO_ANCHOR.MIDDLE; horizontal_rules(c, BORDER)

    note = slide.shapes.add_textbox(Inches(0.2), Inches(6.0), Inches(12.9), Inches(1.2)).text_frame
    note.word_wrap = True
    note.text = ('Experimental public labels are real measurements; predicted sources are in-silico pseudo-labels. '
                 'Per-endpoint source policy is config-vetted (config.RF_SINGLETASK.augmented_sources): '
                 'ADMETlab is dropped for caco2/mdck/ppb (hurts accuracy), solubility uses experimental only. '
                 'Models: single-task RandomForest (H236 features) and multitask Chemprop D-MPNN.')
    note.paragraphs[0].font.size = Pt(10); note.paragraphs[0].font.color.rgb = _hex('#555555')

    dest = OUT / 'adme_datasets.pptx'; prs.save(dest)
    return dest


if __name__ == '__main__':
    dest = build(gather_rows())
    print(f'> wrote {dest}')
