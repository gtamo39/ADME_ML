"""Build a styled potency-summary slide (.pptx) — degradation (DC90/Dmax) + inhibition (IC90/Imax),
color-coded compound columns, left group brackets. Editable native table. Placeholder content lives in
DATA below (edit there to fill real values). Aggregate values only — no SMILES/structures.

Run (env ML):  python python/make_potency_table_slide.py  ->  output/ppt/potency_table.pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from ppt_style import hex_color as _c, horizontal_rules, plain_table

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'output/ppt'; OUT.mkdir(parents=True, exist_ok=True)

NAVY, MUTED = '1F3A6B', '6B7A99'
HEADER_BG, BORDER, MAROON = 'EEF2F8', 'C9D4E5', '7C1D2A'
TITLES = [('title1', 'E65D32'), ('title2', '0EA5CE'), ('title3', '9C3587')]   # ember / azure / purple

SUB = str.maketrans('0123456789maxbn', '₀₁₂₃₄₅₆₇₈₉ₘₐₓᵇₙ')
def _sub(s): return s.translate(SUB)
H_DEG = f'DC{_sub("90")} (nM) / D{_sub("max")} (%)'      # DC₉₀ (nM) / Dₘₐₓ (%)
H_INH = f'IC{_sub("90")} (nM) / I{_sub("max")} (%)'      # IC₉₀ (nM) / Iₘₐₓ (%)

# sections: (section-header, group-label, [(row-label, [(potency, effect) | None per column]) ...])
DATA = [
    (H_DEG, 'Group1', [
        ('Row 4', [('0.030', '97'), ('0.080', '96'), None]),
        ('Row 3', [('0.010', '99'), ('0.013', '100'), None]),
        ('Ro5',   [('1.0', '94'),   ('1.7', '95'),    None]),
    ]),
    (H_INH, 'Group2', [
        ('Row 1', [('0.002', '100'), ('0.002', '100'), ('2.1', '100')]),
        ('Row 2', [('0.009', '97'),  ('0.320*', '97'), ('0.360', '97')]),
    ]),
]


def _style(cell, bg):
    cell.fill.solid(); cell.fill.fore_color.rgb = _c(bg)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.12)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    horizontal_rules(cell, BORDER)                     # horizontal row rules only, no vertical separators


def _text(cell, text, color=NAVY, bold=False, size=12, align=PP_ALIGN.LEFT):
    p = cell.text_frame.paragraphs[0]; p.alignment = align
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run(); r.text = text
    r.font.bold = bold; r.font.size = Pt(size); r.font.color.rgb = _c(color); r.font.name = 'Calibri'


def _value(cell, val):
    p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    if val is None:
        r = p.add_run(); r.text = 'N/A'; r.font.size = Pt(12); r.font.color.rgb = _c(MUTED); r.font.name = 'Calibri'
        return
    potency, effect = val
    r1 = p.add_run(); r1.text = potency; r1.font.bold = True; r1.font.size = Pt(12); r1.font.color.rgb = _c(NAVY); r1.font.name = 'Calibri'
    r2 = p.add_run(); r2.text = f' / {effect}'; r2.font.size = Pt(12); r2.font.color.rgb = _c(NAVY); r2.font.name = 'Calibri'


def build():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    flat = [('header', DATA[0][0])] + [('row', r) for r in DATA[0][2]] \
         + [('header', DATA[1][0])] + [('row', r) for r in DATA[1][2]]
    ncol = 1 + len(TITLES)
    left, top, rowh = Inches(2.15), Inches(1.3), Inches(0.72)
    widths = [3.2, 2.05, 2.05, 2.05]
    tbl = slide.shapes.add_table(len(flat), ncol, left, top, Inches(sum(widths)), rowh * len(flat)).table
    plain_table(tbl)
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i in range(len(flat)):
        tbl.rows[i].height = rowh

    for i, (kind, payload) in enumerate(flat):
        if kind == 'header':
            _style(tbl.cell(i, 0), HEADER_BG); _text(tbl.cell(i, 0), payload, NAVY, bold=True, size=12.5)
            for j, (name, col) in enumerate(TITLES, start=1):
                _style(tbl.cell(i, j), HEADER_BG)
                _text(tbl.cell(i, j), name if payload == H_DEG else '', col, bold=True, size=15, align=PP_ALIGN.CENTER)
        else:
            label, vals = payload
            _style(tbl.cell(i, 0), 'FFFFFF'); _text(tbl.cell(i, 0), label, NAVY, size=12)
            for j, v in enumerate(vals, start=1):
                _style(tbl.cell(i, j), 'FFFFFF'); _value(tbl.cell(i, j), v)

    # left group brackets (maroon bars) + labels, spanning each section's data rows
    r = 0
    for _, group, rows in DATA:
        r += 1                                           # section-header row
        band_top, band_h = top + rowh * r, rowh * len(rows)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left - Inches(0.13), band_top, Inches(0.045), band_h)
        bar.fill.solid(); bar.fill.fore_color.rgb = _c(MAROON); bar.line.fill.background()
        lab = slide.shapes.add_textbox(Inches(0.35), band_top, Inches(1.55), band_h).text_frame
        lab.word_wrap = True; lab.paragraphs[0].text = group
        lab.paragraphs[0].alignment = PP_ALIGN.RIGHT
        lab.vertical_anchor = MSO_ANCHOR.MIDDLE
        f = lab.paragraphs[0].runs[0].font; f.size = Pt(18); f.bold = True; f.color.rgb = _c(NAVY)
        r += len(rows)

    dest = OUT / 'potency_table.pptx'; prs.save(dest)
    return dest


if __name__ == '__main__':
    print('> wrote', build())
